"""Complete the supplied report, preserving its title page and student identity.

Run using the bundled Codex Python. Evidence is actual PTY/Vitest output.
"""

from copy import deepcopy
import hashlib
import json
import re
from pathlib import Path
from lxml import etree

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'output/lab2'
EVIDENCE = OUT / 'evidence'
PICTURES = OUT / 'screenshots'
SOURCE = ROOT / 'Лабораторна_робота_№2_тестування.docx'
DEST = ROOT / 'Лабораторна_робота_№2_тестування_фінальна.docx'

headings = [
    '1. Титульний аркуш, мета роботи, номер варіанта',
    '2. Таблиця аналізу вхідних умов для всіх функцій модуля',
    '3. Лістинги розроблених тестів із поясненням добору вхідних даних',
    '4. Знімки екрана: результат прогону та таблиця покриття',
    '5. Результат експерименту з тестом без твердження та висновок щодо метрики покриття',
    '6. Специфікація, реалізація та тести функції за індивідуальним варіантом',
    '7. Текст рецензії на роботу однокурсника',
    '8. Посилання на репозиторій і запит на злиття',
    '9. Висновки',
]

summaries = {
    s: json.loads((EVIDENCE / f'{s}-summary.json').read_text())
    for s in ('before', 'after', 'final')
}
counts = {
    s: int(re.search(r'Tests\s+(\d+) passed', (EVIDENCE / f'{s}.txt').read_text())[1])
    for s in summaries
}
assert counts == {'before': 145, 'after': 146, 'final': 145}
assert all(x['exitCode'] == 0 for x in json.loads((EVIDENCE / 'execution-status.json').read_text()))
assert 'виклик без перевірки результату' not in (ROOT / 'tests/order-example.test.ts').read_text()
coverage = json.loads((EVIDENCE / 'final-coverage.json').read_text())
order = next(v for v in coverage.values() if v['path'].endswith('/src/order.ts'))
fn = next(v for v in order['fnMap'].values() if v['name'] == 'calcInstallment')
start, end = fn['loc']['start']['line'], fn['loc']['end']['line']
hits = [
    hit for key, b in order['branchMap'].items()
    if b['loc']['start']['line'] >= start and b['loc']['end']['line'] <= end
    for hit in order['b'][key]
]
assert hits and all(h > 0 for h in hits)
order_summary = next(v for k, v in summaries['final'].items() if k.endswith('/src/order.ts'))

doc = Document(SOURCE)
original = Document(SOURCE)
cover = [deepcopy(p._p) for p in original.paragraphs[:20]]
body = doc._element.body
for element in list(body):
    if element.tag != qn('w:sectPr'):
        body.remove(element)
for element in cover:
    body.insert(len(body) - 1, deepcopy(element))

# All cover paragraphs, including the original page break, remain untouched.
# New styles apply only to the completed body of the existing document.
for name in ('ReportBody', 'ReportHeading', 'ReportLabel', 'ReportCode', 'ReportCaption', 'ReportCell'):
    if name not in doc.styles:
        doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    style = doc.styles[name]
    style.base_style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(14)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Times New Roman')
    pf = style.paragraph_format
    pf.line_spacing = 1.5
    pf.space_before = Pt(0)
    pf.space_after = Pt(4)
    pf.first_line_indent = Cm(1.25) if name == 'ReportBody' else Cm(0)
    pf.widow_control = True
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if name == 'ReportBody' else WD_ALIGN_PARAGRAPH.LEFT

doc.styles['ReportHeading'].font.bold = True
doc.styles['ReportHeading'].paragraph_format.keep_with_next = True
doc.styles['ReportHeading'].paragraph_format.keep_together = True
doc.styles['ReportHeading'].paragraph_format.space_before = Pt(10)
outline = OxmlElement('w:outlineLvl')
outline.set(qn('w:val'), '0')
doc.styles['ReportHeading']._element.get_or_add_pPr().append(outline)
doc.styles['ReportLabel'].font.bold = True
doc.styles['ReportLabel'].paragraph_format.keep_with_next = True
doc.styles['ReportLabel'].paragraph_format.space_before = Pt(6)
doc.styles['ReportCode'].paragraph_format.space_after = Pt(0)
doc.styles['ReportCode'].paragraph_format.keep_together = True
doc.styles['ReportCaption'].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.styles['ReportCaption'].paragraph_format.keep_together = True
doc.styles['ReportCell'].paragraph_format.space_after = Pt(0)

assert abs(doc.sections[0].page_width.cm - 21) < 0.01
assert abs(doc.sections[0].page_height.cm - 29.7) < 0.01


def p(text, style='ReportBody'):
    return doc.add_paragraph(text, style)


def heading(number):
    p(headings[number - 1], 'ReportHeading')


def label(text):
    p(text, 'ReportLabel')


def code(text):
    # Whitespace-only source lines are omitted; every nonempty source line is real.
    for line in text.strip('\n').splitlines():
        if line.strip():
            p(line, 'ReportCode')


table_number = 0


def table(title, headers, rows, widths):
    global table_number
    table_number += 1
    caption = p(f'Таблиця {table_number} — {title}', 'ReportCaption')
    caption.paragraph_format.keep_with_next = True
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col, width in zip(t.columns, widths):
        col.width = Cm(width)
    for i, text in enumerate(headers):
        t.rows[0].cells[i].text = text
    for values in rows:
        for cell, text in zip(t.add_row().cells, values):
            cell.text = str(text).replace('_', '_\u200b')
    for row_i, row in enumerate(t.rows):
        trpr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement('w:cantSplit')
        trpr.append(cant_split)
        if row_i == 0:
            repeat = OxmlElement('w:tblHeader')
            trpr.append(repeat)
        for col_i, cell in enumerate(row.cells):
            cell.width = Cm(widths[col_i])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr = cell._tc.get_or_add_tcPr()
            for shd in list(pr.findall(qn('w:shd'))):
                pr.remove(shd)
            margins = OxmlElement('w:tcMar')
            for side, value in [('top', 70), ('bottom', 70), ('left', 70), ('right', 70)]:
                edge = OxmlElement(f'w:{side}')
                edge.set(qn('w:w'), str(value))
                edge.set(qn('w:type'), 'dxa')
                margins.append(edge)
            pr.append(margins)
            for paragraph in cell.paragraphs:
                paragraph.style = doc.styles['ReportCell']
                if row_i == 0:
                    paragraph.paragraph_format.keep_with_next = True
                    for run in paragraph.runs:
                        run.bold = True
    borders = OxmlElement('w:tblBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        edge = OxmlElement(f'w:{side}')
        edge.set(qn('w:val'), 'single')
        edge.set(qn('w:sz'), '6')
        edge.set(qn('w:color'), '000000')
        borders.append(edge)
    t._tbl.tblPr.append(borders)
    return t


figure_number = 0


def figure(filename, caption, width=16.5, crop_bottom=0):
    global figure_number
    figure_number += 1
    paragraph = p('', 'ReportCaption')
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(6)
    shape = paragraph.add_run().add_picture(str(PICTURES / filename), width=Cm(width))
    if crop_bottom:
        picture = shape._inline.graphic.graphicData.pic
        crop = OxmlElement('a:srcRect')
        crop.set('b', str(crop_bottom))
        picture.blipFill.insert(1, crop)
        height = int(shape.height * (1 - crop_bottom / 100000))
        shape.height = height
        picture.spPr.xfrm.ext.cy = height
    shape._inline.docPr.set('descr', caption)
    p(f'Рисунок {figure_number} — {caption}', 'ReportCaption')


heading(1)
for para in original.paragraphs[20:23]:
    p(para.text.strip())
p('Індивідуальне завдання: calcInstallment(total, months) — місячний платіж для 3, 6 та 12 місяців із надбавкою 2 % за кожні три місяці.')
p('Титульний аркуш і дані студентки збережено. Відсутні чотири функції показового прикладу додано за оригінальною методичкою з форматуванням Prettier. Наявну реалізацію варіанта 12 збережено.')
p('Фінальні перевірки виконано 19.10.2026 (Europe/Uzhgorod): 145 тестів у чотирьох файлах. Статус рецензування та Pull Request зазначено в розділах 7 і 8.')

heading(2)
p('Класи еквівалентності об’єднують входи з однаковою очікуваною поведінкою. Межові значення перевіряють зміну правила: сусідні значення до порога, сам поріг і значення після нього. Для дискретного набору строків розстрочки перевірено сусідів кожного допустимого строку.')
table('Аналіз вхідних умов модуля order.ts',
      ['Функція', 'Клас та представники входів', 'Межі й очікувана поведінка'], [
    ['calcSubtotal', 'Непорожній масив, ціна ≥ 0, ціла кількість > 0; кілька позицій.', '100 × 2 + 50 × 3 → 350; враховано кожну позицію.'],
    ['calcSubtotal', 'Нульова ціна; від’ємна ціна; порожній масив.', 'price: −0,01 / 0 / 0,01 → NEGATIVE_PRICE / 0 / 0,01; [] → EMPTY_ORDER.'],
    ['calcSubtotal', 'Кількість: додатна ціла; нуль, від’ємна, дробова, NaN, ±Infinity.', 'qty: 0 / 1 / 2 → INVALID_QTY / 100 / 200 для price = 100; 1,5 також відхиляється.'],
    ['calcDiscountRate', 'basic, silver, gold; сума до порога та від порога.', 'Для 500 → 0 / 0,05 / 0,10; для 2000 → 0,05 / 0,10 / 0,15.'],
    ['calcDiscountRate', 'Поріг великого замовлення; промокод і межа ставки.', '1999 / 2000 / 2001 → 0 / 0,05 / 0,05 для basic; сумарні 0,24 / 0,25 / 0,30 → 0,24 / 0,25 / 0,25.'],
    ['calcDiscountRate', 'Промокод відсутній, нульовий, додатний або від’ємний.', 'Відсутній → 0; −0,01 та −0,1 → INVALID_PROMO; ставка не перевищує 0,25.'],
    ['calcShipping', 'Сума нижча за 1000 та від 1000.', '999 / 1000 / 1001 → 60 / 0 / 0; уточнено межі 999,99 і 1000,01.'],
    ['calcOrderTotal', 'Комбінації позицій, рівня та промокоду; округлення й доставка.', '500 basic → 560; 1200 gold → 1080; 1050 gold → 1005, бо після знижки 945 < 1000.'],
    ['calcOrderTotal', 'Поріг доставки саме після знижки 20 %.', 'Вхід 1249,99 / 1250 / 1250,01 → підсумок 1059,99 / 1000 / 1000,01; рівно 1000 дає безкоштовну доставку.'],
    ['calcOrderTotal', 'Некоректні позиції або промокод.', 'Передає EMPTY_ORDER, NEGATIVE_PRICE, INVALID_QTY та INVALID_PROMO без приховування.'],
    ['calcInstallment', 'Скінченна невід’ємна сума; строки 3, 6, 12.', '1200 → 408 / 208 / 108. Сума −1 / 0 / 1 → INVALID_TOTAL / 0 / 0,34 для 3 місяців.'],
    ['calcInstallment', 'Сусіди допустимих строків; дробовий або непідтримуваний строк.', '2 / 3 / 4; 5 / 6 / 7; 11 / 12 / 13: приймається лише центральне значення. Інші → UNSUPPORTED_MONTHS.'],
    ['calcInstallment', 'Грошове округлення; малі, великі та експоненційні числа.', '0,749 / 0,750 / 0,751 на 3 місяці → 0,25 / 0,26 / 0,26; MIN_VALUE → 0; 10¹⁰ на 12 місяців → 900000000.'],
    ['calcInstallment', 'Від’ємна сума, NaN, ±Infinity; хибні типи під час виконання; переповнення.', 'INVALID_TOTAL або UNSUPPORTED_MONTHS; 10²¹, 1,2 × 10²¹ та MAX_VALUE → INSTALLMENT_OUT_OF_RANGE.'],
], [3.8, 5.7, 7.0])
p('У показовому прикладі збережено саме передбачені перевірки. calcShipping не викидає винятків для некоректних чисел; calcDiscountRate не перевіряє довільний рядок замість Tier, а calcSubtotal не перевіряє скінченність price. Такі додаткові обмеження методичкою не визначені, тому реалізацію прикладу не розширено. Це межі контракту, а не підстава приписувати функціям відсутню валідацію.')

heading(3)
p('Лістинг 1 містить фактичний файл tests/order-example.test.ts. Порожні рядки лістингів опущено для компактності; виконувані рядки та очікувані числа взято з репозиторію. Для варіанта 12 повний набір тестів наведено в розділі 6.')
p('Для calcSubtotal обрано кілька позицій і кількостей, порожній масив, ціну біля нуля та недопустимі кількості. Другу некоректну позицію перевірено окремо, щоб цикл не валідував лише перший елемент. Для calcDiscountRate перевірено всі рівні клієнта, поріг 1999/2000/2001, відсутність промокоду, від’ємний промокод і обмеження ставки 25 %.')
p('Для calcShipping використано обидва класи та сусідів порога, зокрема суми з копійками. Для calcOrderTotal перевірено доставку після знижки, промокод за замовчуванням, кілька позицій, округлення та передавання всіх чотирьох винятків. Випадок 1050 для gold фіксує бізнес-правило методички: знижка переводить 945 грн нижче порога доставки, отже результат становить 1005 грн.')
p('Параметризація it.each застосована до однорідних сценаріїв. Очікувані результати задано числами, незалежними від формули реалізації. toBeCloseTo використано для грошових сум і ставок, toBe — для цілочислового тарифу доставки. У toThrow передається функція, яка викликає модуль, щоб Vitest міг перехопити виняток. Перший тест явно поділено на Arrange — підготовка, Act — виклик, Assert — перевірка. Кожен тест створює власні дані та не залежить від порядку запуску.')
label('Лістинг 1. Тести показового прикладу')
code((ROOT / 'tests/order-example.test.ts').read_text())

heading(4)
p('Нижче наведено нові знімки реального PTY-термінала та згенерованого Vitest HTML-звіту після видалення тимчасового тесту. Термінальний вивід передавався до емулятора xterm без редагування змісту; знімки зроблено браузером автоматично. Первісні знімки з 91 тестом замінено, оскільки вони вже не відповідають завершеному модулю та фінальним 145 тестам.')
figure('tests.png', 'Фінальний npm test: 145 успішних тестів, 4 файли')
p('Збережено всі 79 тестів calcInstallment та 12 тестів попередньої лабораторної роботи: 8 для calculateAverage і 4 для celsiusToFahrenheit. Додано 54 тести показового прикладу. Підсумок узгоджується з кількістю 145 у журналі запуску.')
table('Фінальне покриття коду V8',
      ['Об’єкт', 'Statements', 'Branches', 'Functions', 'Lines'], [
    ['Увесь src', '100 %\n62/62', '100 %\n36/36', '100 %\n12/12', '100 %\n57/57'],
    ['src/order.ts', '100 %\n51/51', '100 %\n28/28', '100 %\n9/9', '100 %\n47/47'],
    ['calcInstallment', 'Окремо не обчислювали', '100 %\n8/8', 'Викликана', 'Рядки виконані'],
], [3.4, 3.25, 3.25, 3.3, 3.3])
figure('final.png', 'Фінальний npm run coverage та окрема перевірка варіанта 12')
figure('index.png', 'HTML Coverage Report із підсумками всіх модулів', crop_bottom=47000)
figure('order-top.png', 'Покриття src/order.ts у фінальному HTML-звіті')
p('HTML-звіт показує дев’ять покритих функцій order.ts з урахуванням допоміжних функцій, конструктора та callback. Відсоток гілок індивідуальної функції обчислено окремо, тому загальне покриття модуля не підміняє перевірку вимоги варіанта.')
figure('order-installment.png', 'Виконані рядки та лічильники calcInstallment у HTML-звіті')
figure('checks.png', 'Успішні ESLint, TypeScript і Prettier')
p('Команди npm test, npm run coverage, npm run lint, npm run typecheck, npm run format:check та npm run coverage:installment завершилися з кодом 0. У coverage-конфігурації збережено пороги: Lines ≥ 80 %, Branches ≥ 80 %, Functions = 100 %. Окремий скрипт додатково вимагає Branches ≥ 90 % саме для calcInstallment.')

heading(5)
p('Експеримент виконано саме для calcOrderTotal, як указано в кроці 6 методички. Спочатку виконано npm run coverage без додаткового тесту. Потім до tests/order-example.test.ts тимчасово додано наведений нижче блок, повторено ту саму команду й збережено обидва JSON-звіти та знімки термінала. Після цього блок видалено, файл відновлено, покриття та інші перевірки виконано повторно.')
label('Лістинг 2. Тимчасовий тест без твердження')
code((EVIDENCE / 'temporary-test.ts.txt').read_text())
table('Порівняння експериментальних вимірювань для всього src',
      ['Показник', 'До тесту', 'Після додавання', 'Після видалення'], [
    ['Тести', '145', '146', '145'],
    ['Statements', '100 % (62/62)', '100 % (62/62)', '100 % (62/62)'],
    ['Branches', '100 % (36/36)', '100 % (36/36)', '100 % (36/36)'],
    ['Functions', '100 % (12/12)', '100 % (12/12)', '100 % (12/12)'],
    ['Lines', '100 % (57/57)', '100 % (57/57)', '100 % (57/57)'],
], [3.6, 4.3, 4.3, 4.3])
p('Зміна кожної з чотирьох метрик становить 0 відсоткових пунктів. Модуль order.ts також залишився на 100 %: Statements 51/51, Branches 28/28, Functions 9/9, Lines 47/47. Додавання тесту не відкрило нових шляхів виконання, оскільки ці шляхи вже були покриті тестами з перевірками.')
figure('before.png', 'Експеримент до додавання тесту: 145 тестів, усі метрики 100 %', crop_bottom=35000)
figure('after.png', 'Експеримент після додавання тесту: 146 тестів, покриття незмінне', crop_bottom=35000)
p('Тест без expect виконує код і виявить необроблений виняток, однак не звіряє повернену суму з очікуваними 155 грн. Якщо функція поверне інше число без винятку, такий тест усе одно пройде. Метрика покриття враховує виконання інструкцій, гілок, функцій і рядків, а не правильність тверджень. За наявності невиконаного шляху порожній виклик міг би підвищити показник без перевірки результату; у цьому експерименті зростання фактично не було.')
p('Отже, 100 % покриття є свідченням виконаних шляхів у цьому наборі даних, але не доказом відсутності дефектів. Якість тестів забезпечують незалежні очікувані значення, змістовні матчери, перевірки винятків та обґрунтовані класи входів. Фінальний набір не містить тимчасового тесту. Для повторення досліду додано команду npm run experiment:order; попередній experiment:installment є лише додатковим дослідом і не підміняє вимогу методички.')

heading(6)
label('Специфікація та прийняті припущення')
p('Явна вимога варіанта 12: calcInstallment(total, months) повертає місячний платіж для 3, 6 та 12 місяців із надбавкою 2 % за кожні три місяці. Загальні вимоги до варіантів передбачають параметризовані тести, перевірку сусідів межі, виняток із власним кодом помилки, явні числові очікування й покриття гілок функції не нижче 90 %.')
p('У наявній реалізації прийнято просту надбавку до початкової суми: 2 % для 3 місяців, 4 % для 6 та 8 % для 12. Розрахунок має зміст total × (1 + 0,02 × months / 3) / months. Для 1200 грн відповідні платежі становлять 408, 208 і 108 грн. Нарахування не є складним відсотком. Цей спосіб тлумачення надбавки збережено як припущення реалізації.')
p('Додаткові припущення: total має бути скінченним числом ≥ 0; нуль дозволений; строк має точно належати множині {3, 6, 12}; результат округлюється до двох десяткових знаків за правилом half up. Округлюється місячний платіж, залишок копійок не переноситься на останній місяць. Зокрема, 0,01 грн на 3 місяці дає 0 грн після округлення. Методичка не задає окремих правил для цих випадків, тому вони не подаються як її дослівні вимоги.')
p('decimalFraction розбирає десятковий запис числа та виконує проміжні дії з bigint, щоб не спотворювати пороги округлення через додаткові двійкові похибки. Вхідний number уже має властиву JavaScript точність; алгоритм працює з його фактичним десятковим записом. При перевищенні безпечного цілого діапазону платежу в копійках реалізація відхиляє розрахунок.')
table('Валідація та коди помилок індивідуальної функції',
      ['Код', 'Причина', 'Приклад перевірки'], [
    ['INVALID_TOTAL', 'Від’ємний або нескінченний total, NaN, хибний тип у runtime.', '−1; NaN; ±Infinity; рядок; null; undefined; {}; bigint.'],
    ['UNSUPPORTED_MONTHS', 'Строк поза {3, 6, 12}; дробовий, нескінченний або хибний тип.', '0; 9; 24; 3,5; NaN; ±Infinity; «3»; null; undefined.'],
    ['INSTALLMENT_OUT_OF_RANGE', 'Округлений платіж у копійках перевищує MAX_SAFE_INTEGER.', '10²¹; 1,2 × 10²¹; Number.MAX_VALUE.'],
], [5.4, 6.1, 5.0])
p('Власний клас InstallmentError розширює Error та має поле code. Якщо обидва аргументи некоректні, спочатку перевіряється total. Це зафіксовано окремим тестом. Реалізацію calcInstallment та її 79 наявних тестів збережено без змін, оскільки перевірка за отриманою методичкою не виявила необхідності їх виправляти.')
label('Лістинг 3. Реалізація варіанта 12 із src/order.ts')
source = (ROOT / 'src/order.ts').read_text()
code(source[source.index('export const INSTALLMENT_MONTHS'):])
label('Пояснення добору тестових даних')
p('Позитивні таблиці містять усі три строки для сум 1200 і 3000 та нуль для кожного строку. Окрема таблиця перевіряє малі дроби, експоненційний запис, MIN_VALUE і велику допустиму суму. Значення 0,749 / 0,750 / 0,751 віддалені на одну тисячну від межі округлення; 0,750 × 1,02 / 3 = 0,255, тому саме середнє значення переходить до 0,26.')
p('Правило «межа мінус один, межа, межа плюс один» застосовано до 3, 6 і 12 місяців та до нижньої межі суми: −1, 0, 1. Набори NaN, нескінченностей і хибних runtime-типів перевіряють відхилення входів, а великі суми перевіряють виняток безпечного діапазону. Очікувані платежі задано літералами. toBeCloseTo застосовано для дробів; toBe — також у збережених тестах з точними еталонними числами та нулем. toThrow(InstallmentError) перевіряє клас, а toThrow(expect.objectContaining(...)) — власний код і, для INVALID_TOTAL, повідомлення.')
label('Лістинг 4. Повний файл tests/order.test.ts')
code((ROOT / 'tests/order.test.ts').read_text())
label('Окрема перевірка покриття гілок')
p(f'Команда npm run coverage:installment читає coverage/coverage-final.json, знаходить calcInstallment у fnMap та відбирає лише branchMap-лічильники всередині її діапазону рядків {start}–{end}. Отримано {len(hits)} покритих лічильників із {len(hits)}, тобто 100 %, що перевищує необхідні 90 %. Загальні Branches 28/28 для order.ts та 36/36 для всіх модулів наведено окремо. Допоміжна decimalFraction знаходиться поза діапазоном calcInstallment, а її гілки враховано в покритті всього модуля.')
figure('installment.png', 'Окрема перевірка calcInstallment: 8/8 гілок, 100 %')
p('Скрипт припиняє виконання з помилкою, якщо файл, функція або лічильники відсутні чи результат нижчий за 90 %. Тому перевірка не зводиться до загального відсотка модуля. Первинні JSON-лічильники та журнал окремої команди збережено в output/lab2/evidence.')

heading(7)
p('Статус: очікує виконання на основі реальної роботи однокурсника. Посилання на Pull Request однокурсника не надано; підтвердженого доступу до такої роботи немає. Рецензію не виконано, коментарі не опубліковано. Вигадані зауваження до невідомого коду не додаються.')
p('Для завершення потрібно отримати доступний URL Pull Request однокурсника, прочитати реалізацію та тести і сформулювати щонайменше три змістовні коментарі з посиланнями на конкретні рядки. Перевіряються наявність межових тестів, незалежність очікувань від продуктової формули та відповідність матчерів типам даних. Також потрібно оцінити назви тестів, зміст тверджень і незалежність тестів. Після фактичного рецензування до цього розділу слід додати справжній текст зауважень та посилання на коментарі.')

heading(8)
repository = p('Репозиторій: ')
link = OxmlElement('w:hyperlink')
link.set(qn('r:id'), doc.part.relate_to('https://github.com/milunchik/testing', RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
run = OxmlElement('w:r')
properties = OxmlElement('w:rPr')
color = OxmlElement('w:color')
color.set(qn('w:val'), '000000')
properties.append(color)
run.append(properties)
value = OxmlElement('w:t')
value.text = 'https://github.com/milunchik/testing'
run.append(value)
link.append(run)
repository._p.append(link)
p('Фактичний origin для fetch і push: git@github.com:milunchik/testing.git. Поточна гілка: feature/lab2-vitest-installment. Початковий HEAD: 7e03d07bdf0c79b4cf3e0838ebfae55b1c1d544f. Перевірену реалізацію та 54 нові тести опубліковано комітом 25fcb2080128cb52a4c9efb701fc729102f71c46; фінальні матеріали звіту додано окремим комітом тієї самої гілки.')
p('Перша перевірка SSH завершилася повідомленням «Permission denied (publickey)». Додаткова перевірка HTTPS підтвердила авторизацію облікового запису milunchik і право push. Гілку опубліковано через HTTPS без зміни origin; через GitHub API створено відкритий Pull Request №2.')
pr = p('Pull Request: ')
link = OxmlElement('w:hyperlink')
link.set(qn('r:id'), doc.part.relate_to('https://github.com/milunchik/testing/pull/2', RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
run = OxmlElement('w:r')
properties = OxmlElement('w:rPr')
color = OxmlElement('w:color')
color.set(qn('w:val'), '000000')
properties.append(color)
run.append(properties)
value = OxmlElement('w:t')
value.text = 'https://github.com/milunchik/testing/pull/2'
run.append(value)
link.append(run)
pr._p.append(link)
p('Базова гілка PR — feature/calculate-average: вона містить попередню лабораторну роботу у власному PR №1. Така база відокремлює зміни лабораторної №2 та зберігає попередні тести. Злиття PR №1 або №2 не виконувалося. Підтверджені GitHub-дані, статус гілки та URL збережено в output/lab2/evidence/github-access.json і pull-request.json. Єдиний незавершений зовнішній етап — рецензія на реальний PR однокурсника.')

heading(9)
p('Під час лабораторної роботи доповнено модуль розрахунку замовлення чотирма функціями показового прикладу відповідно до оригінальної методички та перевірено індивідуальну функцію calcInstallment для варіанта 12. Тестові дані дібрано за класами еквівалентності, межовими значеннями й некоректними входами; використано параметризовані тести, незалежні числові еталони та перевірки винятків. Усі 145 тестів пройшли, зокрема збережені тести попередньої лабораторної роботи. Покриття коду становить 100 % за чотирма метриками; окреме покриття гілок calcInstallment — 100 % (8/8). Перевірки ESLint, TypeScript і Prettier завершилися успішно.').paragraph_format.keep_together = True
p('Експеримент із викликом calcOrderTotal без твердження тимчасово збільшив кількість тестів до 146, але не змінив покриття. Це підтвердило, що виконання коду не перевіряє правильність результату без змістовного твердження. Тимчасовий тест видалено, фінальний стан перевірено повторно. Звіт оновлено реальними доказами виконання, а зміни опубліковано в Pull Request №2. Рецензування однокурсника залишається незавершеним зовнішнім етапом і потребує доступу до його справжнього Pull Request.')

# Remove unused old image relations so stale screenshots do not survive in the package.
used_ids = set(doc._element.xpath('//@r:embed'))
for rid, rel in list(doc.part.rels.items()):
    if rel.reltype.endswith('/image') and rid not in used_ids:
        del doc.part.rels[rid]

actual = [p.text for p in doc.paragraphs if p.style.name == 'ReportHeading']
assert actual == headings
for left, right in zip(cover, list(doc._element.body)[:20]):
    assert etree.tostring(left, method='c14n', exclusive=True) == etree.tostring(right, method='c14n', exclusive=True), 'Original cover was changed'
doc.save(DEST)
manifest = {
    'input': str(SOURCE),
    'output': str(DEST),
    'sections': actual,
    'figures': figure_number,
    'tables': table_number,
    'tests': counts,
    'calcInstallment': {'start': start, 'end': end, 'covered': len(hits), 'total': len(hits), 'percent': 100},
    'original_sha256': hashlib.sha256(SOURCE.read_bytes()).hexdigest(),
    'cover_paragraphs_preserved': 20,
    'peer_review': 'pending: no classmate PR provided',
    'pull_request': 'https://github.com/milunchik/testing/pull/2',
}
(EVIDENCE / 'report-manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + '\n')
print(json.dumps(manifest, ensure_ascii=False, indent=2))
