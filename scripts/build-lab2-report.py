"""Build the Ukrainian report from source listings and measured coverage evidence.

Run with the Codex documents runtime Python, then render with render_docx.py.
"""

import json
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs/lab2"
EVIDENCE = OUTPUT / "evidence"
DOCX = OUTPUT / "Звіт_Лабораторна_2_Варіант_12.docx"


def read(relative):
    return (ROOT / relative).read_text()


def metric(data, name):
    entry = data[name]
    return f"{entry['pct']}% ({entry['covered']}/{entry['total']})"


summaries = {
    stage: json.loads((EVIDENCE / f"{stage}-summary.json").read_text())["total"]
    for stage in ("before", "after", "final")
}
final = summaries["final"]
coverage = json.loads((EVIDENCE / "final-coverage.json").read_text())
order = next(entry for path, entry in coverage.items() if path.endswith("/src/order.ts"))
target = next(entry for entry in order["fnMap"].values() if entry["name"] == "calcInstallment")
start, end = target["loc"]["start"]["line"], target["loc"]["end"]["line"]
branch_ids = [
    key for key, value in order["branchMap"].items()
    if value["loc"]["start"]["line"] >= start and value["loc"]["end"]["line"] <= end
]
hits = [count for key in branch_ids for count in order["b"][key]]
covered = sum(count > 0 for count in hits)
assert hits and covered / len(hits) >= 0.9
counts = {}
for stage in summaries:
    log = (EVIDENCE / f"{stage}.txt").read_text()
    match = re.search(r"Tests\s+(\d+) passed", log)
    assert match, f"Missing actual test count in {stage}"
    counts[stage] = int(match.group(1))
assert not (ROOT / "tests/installment-experiment.test.ts").exists()
verification = (EVIDENCE / "verification.txt").read_text()
assert (EVIDENCE / "verification-status.json").exists()
assert json.loads((EVIDENCE / "verification-status.json").read_text())["exitCode"] == 0


parts = []


def text(value):
    parts.append(value.strip() + "\n")


text("""
# 1 Мета роботи та номер варіанта

Мета роботи — опанувати модульне тестування бізнес-логіки засобами Vitest, застосування матчерових перевірок, параметризованих тестів, аналіз граничних значень та вимірювання покриття коду. Індивідуальний варіант №12 — функція calcInstallment(total, months), яка визначає місячний платіж для строків 3, 6 та 12 місяців із надбавкою 2% за кожні три місяці.

У виконаній частині роботи реалізовано індивідуальну функцію, збережено працездатність лабораторної №1 та отримано 100% покриття наявного коду. Роботу загалом ще не завершено: вихідна методичка лабораторної №2 не була доступна, тому точний приклад із чотирма функціями розрахунку замовлення та передбачений для нього експеримент очікують надання джерела. Зазначені обмеження враховано в усіх висновках цього звіту.

Джерелами фактичних результатів є локальний код, журнали запусків Vitest та JSON-звіти V8. У початковому репозиторії були модулі average.ts і temperature.ts, 12 тестів, конфігурації TypeScript, Vitest, ESLint і Prettier. Залежності та пороги покриття повторно використано без послаблення: рядки не менше 80%, гілки не менше 80%, функції 100%.

# 2 Таблиця аналізу вхідних умов для всіх функцій

Для функцій наданого прикладу нижче наведено лише сценарії, прямо вказані у завданні. Без методички точні ставки, вартість доставки, сигнатури допоміжних типів і очікувані числові результати не визначалися.

| Функція | Вхідні умови та межі | Очікування або стан |
| --- | --- | --- |
| calcSubtotal | Звичайне й порожнє замовлення; від’ємна ціна; некоректна кількість | Очікує точного коду й правил помилок методички |
| calcDiscountRate | basic, silver, gold; промознижка; обмеження; 1999, 2000, 2001 | Очікує ставок, максимальної знижки та правил сумування |
| calcShipping | 999, 1000, 1001; платна й безплатна доставка | Очікує вартості, точного порівняння порогу й залежності від знижки |
| calcOrderTotal | Підсумок; знижка змінює право на безплатну доставку; винятки | Очікує наданої реалізації; поширення винятків ще не перевірено |
| calcInstallment | 1200; строки 3, 6, 12 | Відповідно 408, 208, 108 |
| calcInstallment | Сума −1, 0, 1; строк 3 | INVALID_TOTAL, 0, 0.34 |
| calcInstallment | Строки 2/3/4, 5/6/7, 11/12/13 | Допустимі лише центральні значення; інші відхиляються |
| calcInstallment | NaN, ±Infinity, рядок, null, undefined, об’єкт, bigint | INVALID_TOTAL |
| calcInstallment | Строки −1, 0, 1, 9, 24, 3.5 та нечислові значення | UNSUPPORTED_MONTHS |
| calcInstallment | 0.749, 0.75, 0.751; строк 3 | Відповідно 0.25, 0.26, 0.26 |
| calcInstallment | Дуже малі, експоненційні й надмірно великі значення | Малі платежі округлюються до 0; небезпечний діапазон відхиляється |
| calculateAverage | Додатні, від’ємні, дробові числа; один елемент; нулі; порожній і некоректний масив | Попередні тести та правила EMPTY_ARRAY, INVALID_NUMBER збережено |
| celsiusToFahrenheit | 0, 100, −40 та NaN | 32, 212, −40 та INVALID_TEMPERATURE; попередні тести збережено |

# 3 Лістинги розроблених тестів і вибір тестових даних

Розроблені тести зосереджено у tests/order.test.ts. Усі описи тестів українською мовою, назви функцій і змінних англійською. Використано describe, it, expect, it.each, toBe, toBeCloseTo, toThrow та перевірку властивості code. Перший тест явно поділено на Arrange, Act і Assert; інші короткі випадки зберігають той самий порядок.

Літеральні еталони 408, 208 і 108 перевіряють кожен підтримуваний строк та відповідну надбавку. Друга сума 3000 виключає залежність тестів лише від одного прикладу. Межі строків досліджено за схемою «межа − 1, межа, межа + 1», а для нижньої межі суми обрано −1, 0, 1. Суми 0.749, 0.75, 0.751 перевіряють обидва боки порогу округлення: середній випадок дає рівно 0.255 до округлення.

Дробові значення перевіряються з toBeCloseTo; очікування в таблицях задані незалежними літералами. Вираз 0.1 + 0.2 використано лише як вхід, щоб перевірити типову похибку двійкового подання. Окремі випадки покривають експоненційний запис, найменше додатне число, від’ємний нуль, помилкові типи та пріоритет помилок. Некоректні аргументи інших типів передаються через явне приведення типу тільки в негативних тестах.

Нижче наведено повний лістинг нових тестів без скорочень. Тести чотирьох функцій прикладу ще відсутні через недоступність методички.
""")
text("```typescript\n" + read("tests/order.test.ts").rstrip() + "\n```")
text("# 4 Результати виконання тестів і таблиця покриття коду")
text(f"""
Перед внесенням змін виконано 12 тестів у двох файлах — усі успішно. Після реалізації варіанта №12 та видалення тимчасового тесту виконано {counts['final']} тест у трьох файлах: {counts['final']} успішний, 0 невдалих. Додано 79 тестових випадків. Показники стосуються лише наявного коду і не означають покриття ще не реалізованої частини 1.

| Обсяг | Твердження | Гілки | Функції | Рядки |
| --- | --- | --- | --- | --- |
| Початковий код №1 | 100% (11/11) | 100% (8/8) | 100% (3/3) | 100% (10/10) |
| Підсумок наявного коду | {metric(final, 'statements')} | {metric(final, 'branches')} | {metric(final, 'functions')} | {metric(final, 'lines')} |
| src/order.ts | 100% (25/25) | 100% (12/12) | 100% (3/3) | 100% (25/25) |

Безпосередньо calcInstallment має {covered}/{len(hits)} виконаних лічильників гілок, тобто {covered / len(hits) * 100:.0f}%. Вимогу не менше 90% виконано. Межі функції у виміряному JSON-звіті — рядки {start}–{end}. Перевірка scripts/check-installment-coverage.mjs знаходить функцію у fnMap, відбирає записи branchMap, повністю розташовані в цих межах, та аналізує відповідні масиви b. Отже, результат функції виміряно окремо від файлового показника 12/12.

Виконано npm test, npm run lint, npm run typecheck, npm run format:check, npm run coverage. Усі команди завершилися з кодом 0. Під час розроблення ESLint виявив невизначений глобальний console у допоміжному .mjs-скрипті; проблему усунуто явним імпортом node:console. Правила ESLint не послаблювалися. TypeScript працює в чинному strict-режимі, перевірка формату не переписує попередні матеріали звіту.

Справжні результати збережено в docs/lab2/evidence. Нижче наведено витяг із фактичного журналу запуску покриття після видалення тимчасового тесту.
""")
excerpt = "\n".join(line for line in (EVIDENCE / "final.txt").read_text().splitlines() if re.match(r"\s*(Test Files|Tests|Statements|Branches|Functions|Lines)\s", line))
text("```text\n" + excerpt + "\n```")
text("""
Окремий повний журнал усіх команд перевірки — evidence/verification.txt; код завершення — evidence/verification-status.json. Автоматично отримати знімки екрана не вдалося, оскільки жодного браузера не підключено. За вимогою викладача слід додати справжні знімки npm test, npm run coverage, npm run coverage:installment, перевірок якості та coverage/index.html. Знімки лабораторної №1 для ілюстрації нових результатів не використовувалися.

# 5 Експеримент із тестом без твердження

Точний експеримент методички з calcOrderTotal([item(100)], 'silver') не виконано, оскільки реалізація прикладу недоступна. Щоб дослідити принцип на наявному коді, окремо проведено додатковий експеримент із calcInstallment. Він не підмінює обов’язковий експеримент частини 1.

```typescript
it('виклик без перевірки результату', () => {
  calcInstallment(100, 3);
});
```
""")
experiment_rows = ["| Етап | Тести | Твердження | Гілки | Функції | Рядки |", "| --- | --- | --- | --- | --- | --- |"]
for stage, label in (("before", "До додавання"), ("after", "Із тестом"), ("final", "Після видалення")):
    data = summaries[stage]
    experiment_rows.append(f"| {label} | {counts[stage]} | {metric(data, 'statements')} | {metric(data, 'branches')} | {metric(data, 'functions')} | {metric(data, 'lines')} |")
text("\n".join(experiment_rows))
text("""
Усі чотири показники залишилися 100%; різниця становить 0 відсоткових пунктів. Стверджувати про зростання покриття підстав немає. Додатковий виклик не охопив нових шляхів, оскільки наявні тести вже виконали всі вимірювані ділянки.

Покриття характеризує виконання коду, а не правильність перевірки бізнес-правил. Тест без твердження здатен помітити неочікуваний виняток, але пропустить неправильний числовий результат, якщо функція поверне його без помилки. Високе покриття саме по собі не доводить коректності специфікації, незалежності еталонів чи достатності перевірок. Особливо показово, що 100% покриття наявних файлів не виявляє відсутності обов’язкових функцій прикладу: їх потрібно перевіряти через матрицю вимог.

Скрипт npm run experiment:installment зберігає журнали before.txt, after.txt, final.txt та відповідні JSON-звіти. Власний тимчасовий файл видаляється у finally; повторний запуск без нього знову дав 91 успішний тест. Остаточний набір тестів не містить тесту без твердження.

# 6 Специфікація реалізація та тестування calcInstallment

Доступне формулювання варіанта визначає строки 3, 6, 12 та надбавку 2% за кожні три місяці. Воно не встановлює правила округлення, валідації та капіталізації. Через недоступність повної методички прийнято явні припущення, які необхідно звірити з її текстом до остаточного подання роботи.

Надбавка проста й обчислюється від початкової суми: для трьох місяців 2%, для шести 4%, для дванадцяти 8%. Повна вартість ділиться на кількість місяців. Наприклад, для 1200 та шести місяців загальна сума з надбавкою дорівнює 1248, а місячний платіж — 208. Складні відсотки не застосовуються.

Сума має бути скінченним невід’ємним числом. Нуль дозволено; від’ємний нуль нормалізується до додатного. Дозволено дробові суми з довільною кількістю знаків у наявному поданні number. Округлюється лише кінцевий місячний платіж до двох десяткових знаків за правилом «половину вгору».

Десятковий запис number.toString() перетворюється на раціональний дріб із BigInt. Це дозволяє однозначно округлювати десяткові половини та обробляти експоненційний запис, але не відновлює точність, утрачену до передавання аргументу. Для повернення number платіж у копійках має залишатися в межах безпечних цілих. Різниця від округлення не розподіляється між платежами, оскільки функція не повертає графік погашення.

| Код помилки | Значення |
| --- | --- |
| INVALID_TOTAL | Від’ємна, нескінченна або нечислова сума |
| UNSUPPORTED_MONTHS | Строк не дорівнює числу 3, 6 чи 12 |
| INSTALLMENT_OUT_OF_RANGE | Округлений місячний платіж у копійках перевищує безпечний цілий діапазон |

Клас InstallmentError містить стабільну властивість code та змістовне повідомлення. Якщо хибні обидва аргументи, першою перевіряється сума. Підтримувані строки й параметри надбавки винесено в константи; десяткове перетворення ізольовано у decimalFraction, округлення — в обчисленні roundedCents. За уточнення специфікації ці рішення можна змінити локально.

Нижче наведено повну реалізацію src/order.ts. Тестування описано в розділі 3, а окремий вимір гілок функції — у розділі 4.
""")
text("```typescript\n" + read("src/order.ts").rstrip() + "\n```")
text("""
# 7 Результати взаємного рецензування

Статус реального взаємного рецензування — очікує виконання. Pull Request іншого студента не надано, фактичний код для рецензії не перевірявся, коментарі не публікувалися. Вимогу щонайменше трьох змістовних коментарів не можна вважати виконаною.

Підготовлено окремий шаблон docs/lab2/peer-review-template.md. Перший коментар має зіставляти конкретний фрагмент із бізнес-правилом та незалежним очікуванням. Другий має перевіряти граничні значення, валідацію і коди помилок. Третій має оцінювати якість тверджень, незалежність тестових еталонів і фактичні непокриті шляхи. Для кожного потрібні реальний файл, рядок, спостереження та конкретна пропозиція після перегляду PR.

# 8 Посилання на репозиторій та Pull Request

Адреса наявного origin: https://github.com/milunchik/testing. Вона прочитана з локальної конфігурації Git та не є підтвердженням публікації змін лабораторної №2.

Локальна гілка: feature/lab2-vitest-installment. Її створено від feature/calculate-average, щоб зберегти індивідуальну частину лабораторної №1. Після уточнення користувачки дії з GitHub виключено зі складу роботи; push і створення нового Pull Request не виконувалися. URL нового PR відсутній. Посилання на PR для взаємного рецензування також не надано.

# 9 Висновки

У виконаній частині лабораторної роботи реалізовано та перевірено розрахунок місячного платежу за варіантом №12. Застосовано параметризовані тести, аналіз граничних значень, перевірки винятків і кодів помилок, а також незалежні числові еталони. Наявні модулі лабораторної №1 зберегли працездатність. Повний набір із 91 тесту проходить успішно, а перевірки ESLint, TypeScript і форматування завершуються без помилок.

Окреме вимірювання в межах calcInstallment підтвердило виконання 8 із 8 лічильників гілок, що перевищує мінімальну вимогу 90%. Додатковий експеримент із тестом без твердження не змінив жодного показника покриття. Це демонструє необхідність оцінювати не лише виконані рядки й гілки, а й зміст тверджень, повноту вимог та незалежність очікуваних результатів.

Для остаточного завершення лабораторної роботи необхідно отримати повну методичку, реалізувати її приклад розрахунку замовлення з усіма передбаченими сценаріями та провести точний експеримент із calcOrderTotal. Припущення щодо варіанта №12 слід звірити з вихідним документом. Реальне взаємне рецензування також залишається невиконаним. Отже, високі виміряні показники характеризують якість перевірки виконаної частини й не замінюють завершення решти завдання.
""")

markdown = "\n".join(parts)
(OUTPUT / "report.md").write_text(markdown)
document = Document()
section = document.sections[0]
section.page_width, section.page_height = Cm(21), Cm(29.7)
section.top_margin, section.bottom_margin = Cm(2), Cm(2)
section.left_margin, section.right_margin = Cm(2.5), Cm(1.5)

for style_name in ("Normal", "Title", "Heading 1", "Heading 2", "Caption"):
    style = document.styles[style_name]
    style.font.name = "Times New Roman"
    style.font.size = Pt(14)
    style.font.color.rgb = RGBColor(0, 0, 0)
    properties = style.element.get_or_add_rPr()
    fonts = properties.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        properties.append(fonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{attr}"), "Times New Roman")
    for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        fonts.attrib.pop(qn(f"w:{attr}"), None)
    border = style.element.find(".//" + qn("w:pBdr"))
    if border is not None:
        border.getparent().remove(border)

normal = document.styles["Normal"].paragraph_format
normal.line_spacing = 1.15
normal.space_after = Pt(6)
normal.widow_control = True
normal.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
for name in ("Heading 1", "Heading 2"):
    style = document.styles[name]
    style.font.bold = True
    style.paragraph_format.space_before = Pt(12)
    style.paragraph_format.space_after = Pt(8)
    style.paragraph_format.keep_with_next = True
    style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
document.styles["Title"].paragraph_format.space_after = Pt(10)


def title_paragraph(value, before=0, after=0, bold=False):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.add_run(value).bold = bold
    return paragraph


for value in (
    "МІНІСТЕРСТВО ОСВІТИ І НАУКИ УКРАЇНИ",
    "ДЕРЖАВНИЙ ВИЩИЙ НАВЧАЛЬНИЙ ЗАКЛАД",
    "«УЖГОРОДСЬКИЙ НАЦІОНАЛЬНИЙ УНІВЕРСИТЕТ»",
    "ФАКУЛЬТЕТ ІНФОРМАЦІЙНИХ ТЕХНОЛОГІЙ",
    "КАФЕДРА ПРОГРАМНОГО ЗАБЕЗПЕЧЕННЯ СИСТЕМ",
):
    title_paragraph(value, after=3)
paragraph = document.add_paragraph("Лабораторна робота №2", style="Title")
paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
paragraph.paragraph_format.space_before = Pt(85)
paragraph.runs[0].bold = True
title_paragraph("Модульне тестування бізнес-логіки засобами Vitest\nМатчери, параметризовані тести, покриття коду", after=10)
title_paragraph("Варіант №12", after=0)
paragraph = document.add_paragraph()
paragraph.paragraph_format.space_before = Pt(80)
paragraph.paragraph_format.line_spacing = 1.15
paragraph.paragraph_format.left_indent = Cm(8)
paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
paragraph.add_run("Виконала:\nСтудентка 4 курсу\nДенної форми навчання\nГрупи ІПЗ-1\nГраб Емілія Володимирівна")
title_paragraph("Ужгород – 2026", before=250, bold=True)


def make_table(rows):
    table = document.add_table(rows=1, cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = {2: [6, 11], 3: [4.4, 6.5, 6.1], 5: [4.6, 3.1, 3.1, 3.1, 3.1], 6: [3, 1.8, 3.05, 3.05, 3.05, 3.05]}[len(rows[0])]
    for column, width in zip(table.columns, widths):
        column.width = Cm(width)
    properties = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), "000000")
        borders.append(element)
    properties.append(borders)
    for row_index, cells in enumerate(rows):
        row = table.rows[0] if row_index == 0 else table.add_row()
        no_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(no_split)
        if row_index == 0:
            repeat = OxmlElement("w:tblHeader")
            row._tr.get_or_add_trPr().append(repeat)
        for column_index, (cell, value) in enumerate(zip(row.cells, cells)):
            cell.width = Cm(widths[column_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = value.replace("UNSUPPORTED_MONTHS", "UNSUPPORTED_\nMONTHS").replace("INVALID_TEMPERATURE", "INVALID_\nTEMPERATURE").replace("INSTALLMENT_OUT_OF_RANGE", "INSTALLMENT_\nOUT_OF_RANGE")
            properties = cell._tc.get_or_add_tcPr()
            margins = OxmlElement("w:tcMar")
            for edge in ("top", "left", "bottom", "right"):
                margin = OxmlElement(f"w:{edge}")
                margin.set(qn("w:w"), "80")
                margin.set(qn("w:type"), "dxa")
                margins.append(margin)
            properties.append(margins)
            for p in cell.paragraphs:
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.space_before = Pt(3)
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                if row_index == 0:
                    for run in p.runs:
                        run.bold = True
    document.add_paragraph().paragraph_format.space_after = Pt(0)


lines = markdown.splitlines()
index = 0
in_code = False
while index < len(lines):
    line = lines[index]
    if line.startswith("```"):
        in_code = not in_code
        index += 1
        continue
    if in_code:
        if not line:
            index += 1
            continue
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        paragraph.paragraph_format.widow_control = False
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.add_run(line)
    elif not line.strip():
        pass
    elif line.startswith("# "):
        paragraph = document.add_paragraph(line[2:], style="Heading 1")
        if line[2:3] in {"1", "3"}:
            paragraph.paragraph_format.page_break_before = True
    elif line.startswith("| "):
        rows = []
        while index < len(lines) and lines[index].startswith("| "):
            row = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
            if not all(re.fullmatch(r":?-+:?", cell) for cell in row):
                rows.append(row)
            index += 1
        make_table(rows)
        continue
    else:
        paragraph = document.add_paragraph(line)
        paragraph.paragraph_format.first_line_indent = Cm(1.0)
    index += 1

document.core_properties.title = "Лабораторна робота №2 Модульне тестування засобами Vitest"
document.core_properties.subject = "Варіант 12 calcInstallment"
document.core_properties.author = "Граб Емілія Володимирівна"
document.core_properties.language = "uk-UA"
document.save(DOCX)
reopened = Document(DOCX)
assert len(reopened.tables) == 4
assert all(style.font.name == "Times New Roman" and style.font.size == Pt(14) for style in [reopened.styles[name] for name in ("Normal", "Title", "Heading 1")])
print(DOCX)
